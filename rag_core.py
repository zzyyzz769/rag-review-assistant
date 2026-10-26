"""校务通核心逻辑：本地 RAG、RBAC 权限隔离、审计和问答记录。"""

import csv
import hashlib
import hmac
import io
import json
import re
import secrets
import sqlite3
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from langchain_chroma import Chroma
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_ollama import OllamaEmbeddings, OllamaLLM
from langchain_text_splitters import RecursiveCharacterTextSplitter
from docx import Document as WordDocument
import jieba
from rank_bm25 import BM25Okapi
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}
GROUPS = {
    "student": "学生文档",
    "counselor": "辅导员文档",
    "admin": "管理员文档",
}
ROLE_NAMES = {
    "student": "学生",
    "counselor": "辅导员",
    "admin": "管理员",
}
# 角色可访问的独立知识库。所有检索和原文读取均在此再次校验。
ROLE_ACCESS = {
    "student": ("student",),
    "counselor": ("student", "counselor"),
    "admin": ("student", "counselor", "admin"),
}
COLLECTIONS = {
    "student": "campus_student",
    "counselor": "campus_counselor",
    "admin": "campus_admin",
}
QUERY_EXPANSIONS = {
    "缓考": "缓考 考试 申请 条件 材料 流程 审批",
    "补考": "补考 重修 考试 成绩 申请",
    "重修": "重修 补考 选课 学分 成绩",
    "休学": "休学 复学 学籍 申请 流程",
    "复学": "复学 休学 学籍 申请 流程",
    "退学": "退学 学籍 申请 审批",
    "转专业": "转专业 专业 申请 条件 考核",
    "奖学金": "奖学金 评定 申请 条件 金额",
    "助学金": "助学金 困难认定 申请 材料",
    "请假": "请假 销假 审批 辅导员",
    "选课": "选课 退课 补选 学分",
    "成绩": "成绩 查询 复核 绩点 学分",
}
CORE_SYNONYMS = {
    "缓考": ("缓考", "延期考试", "延缓考试"),
    "补考": ("补考",),
    "重修": ("重修",),
    "休学": ("休学",),
    "复学": ("复学",),
    "退学": ("退学",),
    "转专业": ("转专业", "调整专业"),
    "奖学金": ("奖学金",),
    "助学金": ("助学金",),
    "请假": ("请假",),
    "选课": ("选课",),
    "成绩": ("成绩",),
}
STOPWORDS = {
    "什么", "怎么", "如何", "需要", "可以", "请问", "一下", "知道", "告诉我",
    "本科生", "学生", "学校", "相关", "规定", "办理", "申请", "的", "了", "吗",
}


class CampusRAG:
    """校务通本地核心。actor 必须包含 username 和 role。"""

    def __init__(
        self,
        data_dir: str = "data",
        db_dir: str = "chroma_db",
        system_db: str = "campus_system.db",
        llm_model: str = "qwen2.5:3b",
        embedding_model: str = "bge-m3",
        ollama_base_url: str = "http://localhost:11434",
        chunk_size: int = 500,
        chunk_overlap: int = 80,
        reranker_model: str = "BAAI/bge-reranker-base",
    ) -> None:
        self.data_dir = Path(data_dir)
        self.db_dir = Path(db_dir)
        self.system_db = Path(system_db)
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.db_dir.mkdir(parents=True, exist_ok=True)
        for group in GROUPS:
            (self.data_dir / GROUPS[group]).mkdir(parents=True, exist_ok=True)

        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self._vectorstores: dict[str, Chroma] = {}
        self._chunk_cache: dict[str, list[Document]] = {}
        self._bm25_cache: dict[str, tuple[BM25Okapi, list[list[str]]]] = {}
        self.llm_model = llm_model
        self.reranker_model = reranker_model
        self._reranker: Any = None
        self.embeddings = OllamaEmbeddings(model=embedding_model, base_url=ollama_base_url)
        self.llm = OllamaLLM(model=llm_model, base_url=ollama_base_url, temperature=0.1)
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", "。", "！", "？", "；", ".", " ", ""],
        )
        self._init_system_db()

    # ---------- 通用与权限 ----------

    @staticmethod
    def _now() -> str:
        return datetime.now().astimezone().isoformat(timespec="seconds")

    @staticmethod
    def _role(actor: dict[str, str]) -> str:
        role = str(actor.get("role", ""))
        if role not in ROLE_ACCESS or not actor.get("username"):
            raise PermissionError("登录状态无效，请重新登录。")
        return role

    def _allowed_groups(self, actor: dict[str, str]) -> tuple[str, ...]:
        return ROLE_ACCESS[self._role(actor)]

    def _require_admin(self, actor: dict[str, str]) -> None:
        if self._role(actor) != "admin":
            raise PermissionError("仅管理员可以执行文档管理操作。")

    @staticmethod
    def _safe_filename(filename: str) -> str:
        # 同时处理 Windows 与 POSIX 分隔符，防止路径穿越。
        name = filename.replace("\\", "/").split("/")[-1].strip()
        if not name or name in {".", ".."} or any(ord(c) < 32 for c in name):
            raise ValueError("文件名无效。")
        if Path(name).suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise ValueError("仅支持 PDF、TXT 和 DOCX 文件；旧版 DOC 请先另存为 DOCX。")
        return name

    def _path(self, group: str, filename: str) -> Path:
        if group not in GROUPS:
            raise ValueError("未知文档分组。")
        return self.data_dir / GROUPS[group] / self._safe_filename(filename)

    def _vectorstore(self, group: str) -> Chroma:
        if group not in GROUPS:
            raise ValueError("未知知识库分组。")
        if group not in self._vectorstores:
            self._vectorstores[group] = Chroma(
                collection_name=COLLECTIONS[group],
                embedding_function=self.embeddings,
                persist_directory=str(self.db_dir),
            )
        return self._vectorstores[group]

    # ---------- 用户、密码与审计 ----------

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.system_db, timeout=30)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_system_db(self) -> None:
        with self._connect() as conn:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS users (
                    username TEXT PRIMARY KEY,
                    role TEXT NOT NULL,
                    salt TEXT NOT NULL,
                    password_hash TEXT NOT NULL,
                    iterations INTEGER NOT NULL,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS audit_logs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL,
                    username TEXT NOT NULL,
                    role TEXT NOT NULL,
                    action TEXT NOT NULL,
                    resource TEXT NOT NULL,
                    detail TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS qa_records (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at TEXT NOT NULL,
                    username TEXT NOT NULL,
                    role TEXT NOT NULL,
                    question TEXT NOT NULL,
                    answer TEXT NOT NULL,
                    sources_json TEXT NOT NULL,
                    confidence INTEGER NOT NULL DEFAULT 0,
                    model TEXT NOT NULL DEFAULT '',
                    retrieval_mode TEXT NOT NULL DEFAULT 'hybrid',
                    feedback TEXT NOT NULL DEFAULT '',
                    feedback_at TEXT,
                    rerank_used INTEGER NOT NULL DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS document_checks (
                    group_name TEXT NOT NULL,
                    filename TEXT NOT NULL,
                    sensitive_count INTEGER NOT NULL DEFAULT 0,
                    warning TEXT NOT NULL DEFAULT '',
                    checked_at TEXT NOT NULL,
                    PRIMARY KEY (group_name, filename)
                );
                CREATE INDEX IF NOT EXISTS idx_audit_role
                    ON audit_logs(role, created_at);
                CREATE INDEX IF NOT EXISTS idx_qa_user
                    ON qa_records(username, created_at);
                """
            )

        # 老版本数据库原地升级，不清空历史问答。
        with self._connect() as conn:
            columns = {row["name"] for row in conn.execute("PRAGMA table_info(qa_records)")}
            additions = {
                "confidence": "INTEGER NOT NULL DEFAULT 0",
                "model": "TEXT NOT NULL DEFAULT ''",
                "retrieval_mode": "TEXT NOT NULL DEFAULT 'hybrid'",
                "feedback": "TEXT NOT NULL DEFAULT ''",
                "feedback_at": "TEXT",
                "rerank_used": "INTEGER NOT NULL DEFAULT 0",
            }
            for name, definition in additions.items():
                if name not in columns:
                    conn.execute(f"ALTER TABLE qa_records ADD COLUMN {name} {definition}")

        # 仅首次运行创建演示账户；修改密码后不会被覆盖。
        defaults = (
            ("student", "student", "Student@123"),
            ("counselor", "counselor", "Counselor@123"),
            ("admin", "admin", "Admin@123"),
        )
        for username, role, password in defaults:
            with self._connect() as conn:
                exists = conn.execute(
                    "SELECT 1 FROM users WHERE username = ?", (username,)
                ).fetchone()
            if not exists:
                salt, digest, iterations = self._hash_password(password)
                with self._connect() as conn:
                    conn.execute(
                        "INSERT INTO users VALUES (?, ?, ?, ?, ?, ?)",
                        (username, role, salt, digest, iterations, self._now()),
                    )

    @staticmethod
    def _hash_password(password: str, salt: str | None = None) -> tuple[str, str, int]:
        iterations = 260_000
        salt = salt or secrets.token_hex(16)
        digest = hashlib.pbkdf2_hmac(
            "sha256", password.encode("utf-8"), bytes.fromhex(salt), iterations
        ).hex()
        return salt, digest, iterations

    def _audit(
        self,
        username: str,
        role: str,
        action: str,
        resource: str = "",
        detail: str = "",
    ) -> None:
        with self._connect() as conn:
            conn.execute(
                "INSERT INTO audit_logs (created_at, username, role, action, resource, detail)"
                " VALUES (?, ?, ?, ?, ?, ?)",
                (self._now(), username, role, action, resource, detail[:2000]),
            )

    def authenticate(self, username: str, password: str) -> dict[str, str] | None:
        username = username.strip()
        with self._connect() as conn:
            row = conn.execute(
                "SELECT * FROM users WHERE username = ?", (username,)
            ).fetchone()

        valid = False
        role = "unknown"
        if row:
            role = row["role"]
            # 使用账户保存的迭代数重新计算，兼容未来调整。
            candidate = hashlib.pbkdf2_hmac(
                "sha256",
                password.encode("utf-8"),
                bytes.fromhex(row["salt"]),
                int(row["iterations"]),
            ).hex()
            valid = hmac.compare_digest(candidate, row["password_hash"])

        self._audit(
            username or "空用户名",
            role,
            "login_success" if valid else "login_failed",
            detail="本地账号密码登录",
        )
        return {"username": username, "role": role} if valid else None

    def change_password(
        self, actor: dict[str, str], old_password: str, new_password: str
    ) -> None:
        username = actor["username"]
        self._role(actor)
        if len(new_password) < 8:
            raise ValueError("新密码至少需要 8 个字符。")
        with self._connect() as conn:
            row = conn.execute(
                "SELECT * FROM users WHERE username = ?", (username,)
            ).fetchone()
        if not row:
            raise PermissionError("账户不存在。")
        old_hash = hashlib.pbkdf2_hmac(
            "sha256",
            old_password.encode("utf-8"),
            bytes.fromhex(row["salt"]),
            int(row["iterations"]),
        ).hex()
        if not hmac.compare_digest(old_hash, row["password_hash"]):
            raise ValueError("原密码不正确。")
        salt, digest, iterations = self._hash_password(new_password)
        with self._connect() as conn:
            conn.execute(
                "UPDATE users SET salt = ?, password_hash = ?, iterations = ?"
                " WHERE username = ?",
                (salt, digest, iterations, username),
            )
        self._audit(username, actor["role"], "change_password")

    # ---------- 文档加载与单文档索引 ----------

    @staticmethod
    def _read_txt(path: Path) -> str:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    @staticmethod
    def _read_docx(path: Path) -> str:
        """提取 DOCX 正文与表格；不依赖本机安装 Microsoft Word。"""
        word = WordDocument(str(path))
        lines = [p.text.strip() for p in word.paragraphs if p.text.strip()]
        for table in word.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines).strip()

    @staticmethod
    def _mask_sensitive(text: str) -> str:
        text = re.sub(r"(?<!\d)(1[3-9]\d)\d{4}(\d{4})(?!\d)", r"\1****\2", text)
        text = re.sub(r"(?<!\d)(\d{6})\d{8}(\d{3}[0-9Xx])(?!\d)", r"\1********\2", text)
        text = re.sub(r"((?:学号|学生号)\s*[:：]?\s*)\d{6,20}", r"\1[已脱敏]", text)
        return text

    def preflight_document(
        self, actor: dict[str, str], group: str, filename: str, content: bytes
    ) -> dict[str, Any]:
        """上传前检查可读性、加密、扫描 PDF 与常见敏感信息。"""
        self._require_admin(actor)
        filename = self._safe_filename(filename)
        suffix = Path(filename).suffix.lower()
        warnings: list[str] = []
        encrypted = False
        scanned = False
        fatal = False
        text = ""
        try:
            if suffix == ".pdf":
                reader = PdfReader(io.BytesIO(content))
                encrypted = bool(reader.is_encrypted)
                if encrypted:
                    warnings.append("PDF 已加密，请先解除密码保护")
                else:
                    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
                    text = "\n".join(page_texts)
                    if not reader.pages or len(text.strip()) < max(20, len(reader.pages) * 10):
                        scanned = True
                        warnings.append("疑似扫描/图片型 PDF，当前无法有效提取文字，请先 OCR")
            elif suffix == ".docx":
                if not zipfile.is_zipfile(io.BytesIO(content)):
                    encrypted = True
                    warnings.append("DOCX 无法读取，可能已加密或实际是旧版 DOC")
                else:
                    word = WordDocument(io.BytesIO(content))
                    lines = [p.text.strip() for p in word.paragraphs if p.text.strip()]
                    for table in word.tables:
                        for row in table.rows:
                            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
                    text = "\n".join(lines)
            else:
                raw = content
                for encoding in ("utf-8-sig", "utf-8", "gb18030"):
                    try:
                        text = raw.decode(encoding)
                        break
                    except UnicodeDecodeError:
                        continue
                else:
                    text = raw.decode("utf-8", errors="replace")
        except Exception as exc:
            fatal = True
            warnings.append(f"文件预检失败：{exc}")

        if not encrypted and not scanned and not text.strip():
            fatal = True
            warnings.append("未检测到可读取文字")
        sensitive_counts = {
            "手机号": len(re.findall(r"(?<!\d)1[3-9]\d{9}(?!\d)", text)),
            "身份证号": len(re.findall(r"(?<!\d)\d{17}[0-9Xx](?!\d)", text)),
            "学号": len(re.findall(r"(?:学号|学生号)\s*[:：]?\s*\d{6,20}", text)),
        }
        sensitive_count = sum(sensitive_counts.values())
        if sensitive_count:
            warnings.append(f"检测到 {sensitive_count} 处潜在敏感信息，请确认是否允许入库")
        ok = not fatal and not encrypted and not scanned
        warning = "；".join(warnings)
        with self._connect() as conn:
            conn.execute(
                "INSERT OR REPLACE INTO document_checks"
                " (group_name, filename, sensitive_count, warning, checked_at)"
                " VALUES (?, ?, ?, ?, ?)",
                (group, filename, sensitive_count, warning, self._now()),
            )
        return {
            "ok": ok,
            "warning": warning,
            "encrypted": encrypted,
            "scanned": scanned,
            "sensitive_count": sensitive_count,
            "sensitive_counts": sensitive_counts,
            "char_count": len(text.strip()),
        }

    def _load_file(self, group: str, filename: str) -> list[Document]:
        path = self._path(group, filename)
        if not path.is_file():
            raise FileNotFoundError(f"文件不存在：{GROUPS[group]}/{filename}")
        documents: list[Document] = []
        if path.suffix.lower() == ".pdf":
            for page in PyPDFLoader(str(path)).load():
                text = page.page_content.strip()
                if text:
                    documents.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": filename,
                                "page": int(page.metadata.get("page", 0)) + 1,
                                "file_type": "pdf",
                            },
                        )
                    )
        elif path.suffix.lower() == ".docx":
            text = self._read_docx(path)
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": filename, "file_type": "docx"},
                    )
                )
        else:
            text = self._read_txt(path).strip()
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": filename, "file_type": "txt"},
                    )
                )
        if not documents:
            raise ValueError(
                f"{filename} 中没有可读取的文字；扫描版 PDF 或纯图片 Word 需先 OCR。"
            )
        return documents

    def save_document(
        self, actor: dict[str, str], group: str, filename: str, content: bytes
    ) -> str:
        self._require_admin(actor)
        if not content:
            raise ValueError(f"{filename} 是空文件。")
        path = self._path(group, filename)
        path.write_bytes(content)
        self._audit(
            actor["username"], actor["role"], "upload_document",
            f"{group}/{path.name}", f"{len(content)} bytes",
        )
        return path.name

    def index_document(self, actor: dict[str, str], group: str, filename: str) -> int:
        """只删除并重写当前文件的向量，不影响同库其他文档。"""
        self._require_admin(actor)
        filename = self._safe_filename(filename)
        documents = self._load_file(group, filename)
        chunks = [d for d in self.splitter.split_documents(documents) if d.page_content.strip()]
        for index, doc in enumerate(chunks, start=1):
            doc.metadata.update({
                "group": group,
                "source": filename,
                "doc_id": f"{group}/{filename}",
                "chunk": index,
            })
        store = self._vectorstore(group)
        store.delete(where={"source": filename})
        if chunks:
            for start in range(0, len(chunks), 32):
                store.add_documents(chunks[start:start + 32])
        self._chunk_cache.pop(group, None)
        self._bm25_cache.pop(group, None)
        self._audit(
            actor["username"], actor["role"], "index_document",
            f"{group}/{filename}", f"{len(chunks)} 个切片",
        )
        return len(chunks)

    def delete_document(self, actor: dict[str, str], group: str, filename: str) -> None:
        self._require_admin(actor)
        path = self._path(group, filename)
        self._vectorstore(group).delete(where={"source": path.name})
        path.unlink(missing_ok=True)
        self._chunk_cache.pop(group, None)
        self._bm25_cache.pop(group, None)
        with self._connect() as conn:
            conn.execute(
                "DELETE FROM document_checks WHERE group_name = ? AND filename = ?",
                (group, path.name),
            )
        self._audit(
            actor["username"], actor["role"], "delete_document", f"{group}/{path.name}"
        )

    def list_documents(self, actor: dict[str, str]) -> list[dict[str, Any]]:
        result = []
        for group in self._allowed_groups(actor):
            folder = self.data_dir / GROUPS[group]
            for path in sorted(folder.iterdir()):
                if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                    result.append({
                        "doc_id": f"{group}/{path.name}",
                        "group": group,
                        "group_name": GROUPS[group],
                        "filename": path.name,
                        "size_kb": round(path.stat().st_size / 1024, 1),
                        "modified": datetime.fromtimestamp(path.stat().st_mtime).strftime(
                            "%Y-%m-%d %H:%M:%S"
                        ),
                    })
        return result

    def accessible_chunk_count(self, actor: dict[str, str]) -> int:
        return sum(
            int(self._vectorstore(group)._collection.count())
            for group in self._allowed_groups(actor)
        )

    def get_stats(self, actor: dict[str, str]) -> dict[str, int]:
        groups = self._allowed_groups(actor)
        placeholders = ",".join("?" for _ in groups)
        with self._connect() as conn:
            sensitive = conn.execute(
                f"SELECT COALESCE(SUM(sensitive_count), 0) FROM document_checks"
                f" WHERE group_name IN ({placeholders})",
                groups,
            ).fetchone()[0]
            today = datetime.now().astimezone().date().isoformat()
            today_qa = conn.execute(
                "SELECT COUNT(*) FROM qa_records WHERE username = ? AND created_at >= ?",
                (actor["username"], today),
            ).fetchone()[0]
        return {
            "documents": len(self.list_documents(actor)),
            "chunks": self.accessible_chunk_count(actor),
            "sensitive_hits": int(sensitive),
            "today_qa": int(today_qa),
        }

    # ---------- 检索、问答与原文 ----------

    @staticmethod
    def _tokens(text: str) -> list[str]:
        # 修复部分 PDF 在汉字之间插入空格导致的分词断裂。
        text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
        tokens = []
        for token in jieba.lcut(text.lower()):
            token = token.strip()
            if len(token) >= 2 and token not in STOPWORDS and re.search(r"\w", token):
                tokens.append(token)
        return tokens

    @staticmethod
    def _expanded_query(question: str) -> str:
        additions = [value for key, value in QUERY_EXPANSIONS.items() if key in question]
        return question + (" " + " ".join(additions) if additions else "")

    def _group_chunks(self, group: str) -> list[Document]:
        """读取一个权限组的切片用于 BM25；索引变更时会主动失效缓存。"""
        if group not in self._chunk_cache:
            result = self._vectorstore(group).get(include=["documents", "metadatas"])
            documents = result.get("documents") or []
            metadatas = result.get("metadatas") or []
            self._chunk_cache[group] = [
                Document(page_content=text, metadata=metadata or {})
                for text, metadata in zip(documents, metadatas)
                if text and (metadata or {}).get("group") == group
            ]
        return self._chunk_cache[group]

    def _keyword_candidates(
        self, group: str, question: str, limit: int
    ) -> list[tuple[Document, float]]:
        docs = self._group_chunks(group)
        if not docs:
            return []
        if group not in self._bm25_cache:
            tokenized = [self._tokens(doc.page_content) for doc in docs]
            if not any(tokenized):
                return []
            self._bm25_cache[group] = (BM25Okapi(tokenized), tokenized)
        bm25, _ = self._bm25_cache[group]
        query_tokens = self._tokens(question)
        if not query_tokens:
            return []
        scores = bm25.get_scores(query_tokens)
        ranked = sorted(enumerate(scores), key=lambda item: item[1], reverse=True)
        return [
            (docs[index], float(score))
            for index, score in ranked[:limit]
            if score > 0
        ]

    @staticmethod
    def _doc_key(doc: Document) -> str:
        metadata = doc.metadata
        return (
            f"{metadata.get('group')}/{metadata.get('source')}/"
            f"{metadata.get('page')}/{metadata.get('chunk')}"
        )

    def _evidence(self, doc: Document, question: str) -> tuple[str, str, list[str]]:
        """从切片中抽取最关键的 1～2 句，而不是把整段交给用户阅读。"""
        text = re.sub(r"[ \t]+", " ", doc.page_content.strip())
        sentences = [
            part.strip()
            for part in re.split(r"(?<=[。！？；])|\n+", text)
            if part.strip()
        ]
        if not sentences:
            sentences = [text]
        query_tokens = set(self._tokens(question))
        core_terms = [key for key in QUERY_EXPANSIONS if key in question]
        search_terms = {
            synonym
            for term in core_terms
            for synonym in CORE_SYNONYMS.get(term, (term,))
        }

        def score(sentence: str) -> int:
            compact = re.sub(r"\s+", "", sentence)
            value = sum(3 if len(token) > 2 else 2 for token in query_tokens if token in compact)
            value += sum(8 for term in search_terms if term in compact)
            return value

        best_index = max(range(len(sentences)), key=lambda i: score(sentences[i]))
        chosen = {best_index}
        if len(sentences) > 1:
            second_index = max(
                (i for i in range(len(sentences)) if i != best_index),
                key=lambda i: score(sentences[i]),
            )
            if score(sentences[second_index]) > 0:
                chosen.add(second_index)
        selected = [sentences[i] for i in sorted(chosen)]
        evidence = " …… ".join(selected)
        if len(evidence) > 360:
            evidence = sentences[best_index][:357] + "……"
        compact_text = re.sub(r"\s+", "", text)
        matched = sorted(
            {token for token in query_tokens if token in compact_text} |
            {term for term in search_terms if term in compact_text}
        )
        return evidence, sentences[best_index], matched[:6]

    def _rerank(
        self, question: str, entries: list[dict[str, Any]]
    ) -> list[dict[str, Any]]:
        """可选的本地 CrossEncoder 重排，首次使用需下载模型。"""
        if not entries:
            return entries
        if self._reranker is None:
            try:
                from sentence_transformers import CrossEncoder
            except ImportError as exc:
                raise RuntimeError(
                    "启用 Rerank 需要安装 sentence-transformers。"
                ) from exc
            self._reranker = CrossEncoder(self.reranker_model, device="cpu")
        pairs = [[question, entry["doc"].page_content] for entry in entries]
        scores = self._reranker.predict(pairs)
        for entry, score in zip(entries, scores):
            entry["rerank_score"] = float(score)
        return sorted(entries, key=lambda item: item["rerank_score"], reverse=True)

    def ask(
        self,
        actor: dict[str, str],
        question: str,
        top_k: int = 4,
        retrieval_mode: str = "hybrid",
        use_rerank: bool = False,
    ) -> dict[str, Any]:
        allowed = self._allowed_groups(actor)
        question = question.strip()
        if not question:
            raise ValueError("问题不能为空。")
        if retrieval_mode not in {"vector", "bm25", "hybrid"}:
            raise ValueError("未知检索模式。")
        top_k = max(1, min(int(top_k), 6))
        candidate_limit = max(12, top_k * 4)
        merged: dict[str, dict[str, Any]] = {}

        def add_candidate(doc: Document, rank: int, channel: str, score: float) -> None:
            key = self._doc_key(doc)
            entry = merged.setdefault(key, {
                "doc": doc, "rrf": 0.0, "keyword_score": 0.0,
                "vector_distance": None, "rerank_score": None,
            })
            weight = 1.35 if channel == "keyword" else 1.0
            entry["rrf"] += weight / (60 + rank + 1)
            if channel == "keyword":
                entry["keyword_score"] = max(entry["keyword_score"], score)
            else:
                entry["vector_distance"] = score

        expanded_query = self._expanded_query(question)
        for group in allowed:
            if retrieval_mode in {"vector", "hybrid"}:
                vector_results = self._vectorstore(group).similarity_search_with_score(
                    expanded_query, k=candidate_limit
                )
                for rank, (doc, distance) in enumerate(vector_results):
                    if doc.metadata.get("group") == group:
                        add_candidate(doc, rank, "vector", float(distance))
            if retrieval_mode in {"bm25", "hybrid"}:
                for rank, (doc, score) in enumerate(
                    self._keyword_candidates(group, question, candidate_limit)
                ):
                    add_candidate(doc, rank, "keyword", score)

        candidates = list(merged.values())
        core_terms = [key for key in QUERY_EXPANSIONS if key in question]
        if core_terms:
            search_terms = {
                synonym for term in core_terms
                for synonym in CORE_SYNONYMS.get(term, (term,))
            }
            if retrieval_mode == "hybrid":
                # 混合模式补充全库直接核查，避免双路召回共同漏掉相关页。
                direct_entries: dict[str, dict[str, Any]] = {}
                for group in allowed:
                    for doc in self._group_chunks(group):
                        compact = re.sub(r"\s+", "", doc.page_content)
                        if not any(term in compact for term in search_terms):
                            continue
                        key = self._doc_key(doc)
                        entry = merged.setdefault(key, {
                            "doc": doc, "rrf": 0.0, "keyword_score": 0.0,
                            "vector_distance": None, "rerank_score": None,
                        })
                        entry["rrf"] += 0.04
                        direct_entries[key] = entry
                candidates = list(direct_entries.values())
            else:
                candidates = [
                    entry for entry in candidates
                    if any(term in re.sub(r"\s+", "", entry["doc"].page_content)
                           for term in search_terms)
                ]
        elif retrieval_mode != "vector":
            keyword_hits = [entry for entry in candidates if entry["keyword_score"] > 0]
            if keyword_hits:
                candidates = keyword_hits

        candidates.sort(key=lambda entry: entry["rrf"], reverse=True)
        if use_rerank and candidates:
            candidates = self._rerank(question, candidates[:30])

        max_bm25 = max((entry["keyword_score"] for entry in candidates), default=0.0)
        selected_entries = []
        seen_pages = set()
        for entry in candidates:
            doc = entry["doc"]
            page_key = (doc.metadata.get("group"), doc.metadata.get("source"),
                        doc.metadata.get("page"))
            if page_key in seen_pages:
                continue
            seen_pages.add(page_key)
            bm25_norm = entry["keyword_score"] / max_bm25 if max_bm25 > 0 else 0.0
            distance = entry["vector_distance"]
            vector_similarity = 1.0 / (1.0 + max(distance, 0.0)) if distance is not None else 0.0
            if retrieval_mode == "bm25":
                source_confidence = 100.0 * bm25_norm
            elif retrieval_mode == "vector":
                source_confidence = 100.0 * vector_similarity
            else:
                source_confidence = 100.0 * (0.55 * bm25_norm + 0.45 * vector_similarity)
            entry["confidence"] = max(0, min(100, round(source_confidence)))
            selected_entries.append(entry)
            if len(selected_entries) >= top_k:
                break

        if not selected_entries:
            topic = "、".join(core_terms) if core_terms else question
            answer = (
                f"未在当前权限范围内的文档中找到与“{topic}”直接相关的规定。"
                "为避免误导，本次不附无关来源。请补充相关制度或换用原文表述。"
            )
            sources: list[dict[str, Any]] = []
            confidence = 0
        else:
            blocks = []
            sources = []
            for index, entry in enumerate(selected_entries, start=1):
                doc = entry["doc"]
                page = doc.metadata.get("page")
                location = f"，第 {page} 页" if page else ""
                group = str(doc.metadata["group"])
                evidence, highlight, matched = self._evidence(doc, question)
                label = f"{GROUPS[group]}/{doc.metadata['source']}{location}"
                blocks.append(f"[资料{index}] 来源：{label}\n{doc.page_content.strip()}")
                sources.append({
                    "id": index, "group": group, "group_name": GROUPS[group],
                    "source": doc.metadata["source"], "page": page,
                    "chunk": doc.metadata.get("chunk"), "content": doc.page_content.strip(),
                    "evidence": evidence, "highlight": highlight, "matched": matched,
                    "confidence": entry["confidence"],
                })
            prompt = f"""你是“校务通”校园教务问答助手。目标是直接解决用户问题，而不是堆砌资料。
1. 只使用与问题直接相关的资料，忽略其他资料，不得用相近制度替代原问题。
2. 第一行先给明确结论；随后用不超过4条短句列出条件、材料或步骤，并标注[资料1]等编号。
3. 如果资料只提到该事项但没有所需条件或流程，直接说明缺少哪类信息。
4. 不得编造条款、日期、部门、材料或流程，不复述这些指令。

用户问题：{question}

资料：
{chr(10).join(blocks)}
"""
            answer = str(self.llm.invoke(prompt)).strip()
            values = [entry["confidence"] for entry in selected_entries]
            confidence = round(0.6 * max(values) + 0.4 * sum(values) / len(values))

        sources_json = json.dumps(sources, ensure_ascii=False)
        with self._connect() as conn:
            cursor = conn.execute(
                "INSERT INTO qa_records (created_at, username, role, question, answer,"
                " sources_json, confidence, model, retrieval_mode, rerank_used)"
                " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (self._now(), actor["username"], actor["role"], question, answer,
                 sources_json, confidence, self.llm_model, retrieval_mode, int(use_rerank)),
            )
            record_id = int(cursor.lastrowid)
        self._audit(
            actor["username"], actor["role"], "query", ",".join(allowed),
            f"问题：{question[:500]}；模式：{retrieval_mode}；Rerank：{use_rerank}；"
            f"置信度：{confidence}；命中："
            + "；".join(f"{s['group']}/{s['source']}" for s in sources),
        )
        return {
            "answer": answer, "sources": sources, "record_id": record_id,
            "confidence": confidence, "model": self.llm_model,
            "retrieval_mode": retrieval_mode, "rerank_used": use_rerank,
        }

    def open_source(
        self,
        actor: dict[str, str],
        group: str,
        filename: str,
        page: int | None = None,
    ) -> dict[str, str]:
        if group not in self._allowed_groups(actor):
            self._audit(actor.get("username", "?"), actor.get("role", "?"),
                        "access_denied", f"{group}/{filename}", "越权读取原文")
            raise PermissionError("当前角色无权查看该来源文档。")
        documents = self._load_file(group, filename)
        if page is not None:
            pages = [d for d in documents if d.metadata.get("page") == page]
            if not pages:
                raise ValueError("未找到对应 PDF 页面。")
            text = "\n\n".join(d.page_content for d in pages)
            location = f"第 {page} 页"
        else:
            text = "\n\n".join(d.page_content for d in documents)
            location = "DOCX 全文" if Path(filename).suffix.lower() == ".docx" else "文档全文"
        self._audit(
            actor["username"], actor["role"], "view_source",
            f"{group}/{filename}", location,
        )
        return {"text": text, "location": location}

    # ---------- 历史、CSV 与审计查询 ----------

    def get_qa_records(self, actor: dict[str, str], limit: int = 30) -> list[dict[str, Any]]:
        self._role(actor)
        with self._connect() as conn:
            rows = conn.execute(
                "SELECT * FROM qa_records WHERE username = ? ORDER BY id DESC LIMIT ?",
                (actor["username"], max(1, min(limit, 200))),
            ).fetchall()
        records = []
        for row in rows:
            records.append({
                "id": row["id"],
                "created_at": row["created_at"],
                "question": row["question"],
                "answer": row["answer"],
                "sources": json.loads(row["sources_json"]),
                "confidence": int(row["confidence"]),
                "model": row["model"],
                "retrieval_mode": row["retrieval_mode"],
                "rerank_used": bool(row["rerank_used"]),
                "feedback": row["feedback"],
                "feedback_at": row["feedback_at"],
            })
        return records

    def set_feedback(self, actor: dict[str, str], record_id: int, feedback: str) -> None:
        self._role(actor)
        if feedback not in {"useful", "useless"}:
            raise ValueError("未知反馈类型。")
        with self._connect() as conn:
            cursor = conn.execute(
                "UPDATE qa_records SET feedback = ?, feedback_at = ?"
                " WHERE id = ? AND username = ?",
                (feedback, self._now(), record_id, actor["username"]),
            )
        if cursor.rowcount == 0:
            raise PermissionError("只能评价自己的问答记录。")
        self._audit(
            actor["username"], actor["role"], "answer_feedback",
            str(record_id), "有用" if feedback == "useful" else "无用",
        )

    @staticmethod
    def _csv_safe(value: Any) -> str:
        """避免用户文本在 Excel 中被当作公式执行。"""
        text = str(value)
        return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text

    @staticmethod
    def _watermark(username: str) -> str:
        """将当前用户 ID 编码为零宽字符，附加到导出内容中用于溯源。"""
        bits = "".join(f"{byte:08b}" for byte in username.encode("utf-8"))
        return "\u2063" + "".join("\u200b" if bit == "0" else "\u200c" for bit in bits) + "\u2063"

    def export_qa_csv(self, actor: dict[str, str]) -> bytes:
        self._role(actor)
        with self._connect() as conn:
            rows = conn.execute(
                "SELECT * FROM qa_records WHERE username = ? ORDER BY id ASC",
                (actor["username"],),
            ).fetchall()
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow([
            "时间", "用户名", "角色", "问题", "回答", "置信度",
            "模型", "检索模式", "Rerank", "反馈", "原文来源",
        ])
        for row in rows:
            sources = json.loads(row["sources_json"])
            source_text = "\n".join(
                f"[资料{s['id']}] {s['group_name']}/{s['source']}"
                + (f" 第{s['page']}页" if s.get("page") else "")
                + f" 切片{s.get('chunk', '')}\n{s.get('evidence', s['content'])}"
                for s in sources
            )
            writer.writerow([
                self._csv_safe(value) for value in (
                    row["created_at"], row["username"],
                    ROLE_NAMES.get(row["role"], row["role"]),
                    row["question"], row["answer"] + self._watermark(actor["username"]),
                    row["confidence"], row["model"], row["retrieval_mode"],
                    "是" if row["rerank_used"] else "否", row["feedback"], source_text,
                )
            ])
        self._audit(actor["username"], actor["role"], "export_qa_csv", detail=f"{len(rows)} 条")
        # UTF-8 BOM，方便 Windows Excel 直接打开而不乱码。
        return output.getvalue().encode("utf-8-sig")

    def get_audit_logs(
        self, actor: dict[str, str], role: str | None = None, limit: int = 200
    ) -> list[dict[str, Any]]:
        self._require_admin(actor)
        limit = max(1, min(int(limit), 1000))
        sql = "SELECT * FROM audit_logs"
        params: list[Any] = []
        if role in ROLE_NAMES:
            sql += " WHERE role = ?"
            params.append(role)
        sql += " ORDER BY id DESC LIMIT ?"
        params.append(limit)
        with self._connect() as conn:
            rows = conn.execute(sql, params).fetchall()
        return [dict(row) for row in rows]


    @staticmethod
    def _mask_username(username: str) -> str:
        if not username:
            return "***"
        if len(username) == 1:
            return username[0] + "***"
        return username[0] + "***" + username[-1]

    def export_audit_csv(self, actor: dict[str, str]) -> bytes:
        """管理员导出脱敏审计日志，并嵌入当前管理员 ID 水印。"""
        self._require_admin(actor)
        rows = list(reversed(self.get_audit_logs(actor, limit=1000)))
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow([
            "脱敏标记", "访问时间", "用户ID（脱敏）", "访问角色",
            "操作", "命中文档/资源", "详情（脱敏）",
        ])
        for row in rows:
            detail = self._mask_sensitive(row["detail"]) + self._watermark(actor["username"])
            writer.writerow([
                self._csv_safe(value) for value in (
                    "已脱敏", row["created_at"], self._mask_username(row["username"]),
                    ROLE_NAMES.get(row["role"], row["role"]), row["action"],
                    self._mask_sensitive(row["resource"]), detail,
                )
            ])
        self._audit(
            actor["username"], actor["role"], "export_audit_csv",
            detail=f"导出 {len(rows)} 条，已脱敏并嵌入水印",
        )
        return output.getvalue().encode("utf-8-sig")